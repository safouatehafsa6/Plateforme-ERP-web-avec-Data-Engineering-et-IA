"""Traitement intelligent de documents (IDP) - étape 1 : lecture des formats.

Entrée  : les octets d'un fichier + son nom.
Sortie  : un `DocumentLu` (texte brut + tableaux) quel que soit le format :
          PDF (texte ou numérisé), images, Word, Excel, CSV, PowerPoint,
          XML, JSON, texte.

Principes de sécurité :
- le format est déterminé par le CONTENU (signature du fichier), jamais par
  la seule extension, qui peut mentir ;
- taille maximale du fichier, et protection contre les archives Office
  piégées (bombes de décompression) ;
- XML lu avec defusedxml (refuse les entités externes : attaque XXE) ;
- les formats Office anciens (.doc, .xls, .ppt) sont refusés avec un message
  clair plutôt que mal lus.

Les bibliothèques lourdes sont importées à la demande : ce module reste
importable même si l'une d'elles manque, et l'erreur n'apparaît que pour le
format concerné.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import zipfile
from dataclasses import dataclass, field

TAILLE_MAX_OCTETS = 20 * 1024 * 1024            # 20 Mo par fichier
TAILLE_MAX_DECOMPRESSEE = 200 * 1024 * 1024     # protection bombe zip
NB_MAX_ENTREES_ZIP = 5000
MAX_PAGES_OCR = 15
MAX_LIGNES_TABLEUR = 5000
MAX_CARACTERES = 200_000
MAX_PIXELS_IMAGE = 50_000_000
SEUIL_TEXTE_PAR_PAGE = 40                       # en dessous : PDF considéré numérisé
LANGUES_OCR = "fra+ara+eng"


class DocumentIllisible(Exception):
    """Le fichier ne peut pas être lu (message destiné à l'utilisateur)."""


class FormatNonSupporte(DocumentIllisible):
    pass


class FichierTropVolumineux(DocumentIllisible):
    pass


@dataclass
class DocumentLu:
    format: str
    texte: str = ""
    tableaux: list = field(default_factory=list)   # liste de tableaux : [lignes][cellules]
    pages: int | None = None
    ocr_utilise: bool = False
    avertissements: list = field(default_factory=list)


def empreinte_fichier(contenu: bytes) -> str:
    """SHA-256 du fichier : sert à détecter qu'un même fichier est importé deux fois."""
    return hashlib.sha256(contenu).hexdigest()


# --------------------------------------------------------------------------
# Détection du format (par le contenu)
# --------------------------------------------------------------------------

_SIGNATURES_IMAGE = (
    b"\x89PNG\r\n\x1a\n", b"\xff\xd8\xff", b"GIF87a", b"GIF89a",
    b"II*\x00", b"MM\x00*", b"BM",
)
_SIGNATURE_OLE = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"   # .doc / .xls / .ppt anciens


def _extension(nom: str) -> str:
    return nom.lower().rsplit(".", 1)[-1] if "." in nom else ""


def decoder_texte(contenu: bytes) -> str:
    if contenu.startswith(b"\xff\xfe") or contenu.startswith(b"\xfe\xff"):
        return contenu.decode("utf-16", errors="replace")
    for encodage in ("utf-8-sig", "cp1252"):
        try:
            return contenu.decode(encodage)
        except UnicodeDecodeError:
            continue
    return contenu.decode("latin-1", errors="replace")


def _verifier_zip(contenu: bytes) -> zipfile.ZipFile:
    try:
        archive = zipfile.ZipFile(io.BytesIO(contenu))
    except zipfile.BadZipFile as exc:
        raise DocumentIllisible("Fichier Office corrompu ou illisible.") from exc
    infos = archive.infolist()
    if len(infos) > NB_MAX_ENTREES_ZIP or sum(i.file_size for i in infos) > TAILLE_MAX_DECOMPRESSEE:
        raise FichierTropVolumineux("Fichier Office refusé : contenu décompressé anormalement volumineux.")
    return archive


def detecter_format(contenu: bytes, nom_fichier: str = "") -> str:
    if not contenu:
        raise DocumentIllisible("Le fichier est vide.")

    if contenu.startswith(b"%PDF"):
        return "pdf"
    if contenu.startswith(_SIGNATURE_OLE):
        raise FormatNonSupporte(
            "Les anciens formats Office (.doc, .xls, .ppt) ne sont pas pris en charge. "
            "Enregistrez le fichier en .docx, .xlsx ou .pptx."
        )
    if contenu[:4] == b"RIFF" and contenu[8:12] == b"WEBP":
        return "image"
    if any(contenu.startswith(sig) for sig in _SIGNATURES_IMAGE):
        return "image"

    if contenu.startswith(b"PK\x03\x04"):
        noms = set(_verifier_zip(contenu).namelist())
        if "word/document.xml" in noms:
            return "docx"
        if "xl/workbook.xml" in noms:
            return "xlsx"
        if "ppt/presentation.xml" in noms:
            return "pptx"
        raise FormatNonSupporte("Archive non reconnue : seuls Word, Excel et PowerPoint (.docx, .xlsx, .pptx) sont acceptés.")

    # Formats texte : l'extension sert d'indice, le contenu tranche.
    if b"\x00" in contenu[:2048] and not contenu.startswith((b"\xff\xfe", b"\xfe\xff")):
        raise FormatNonSupporte("Format de fichier non reconnu.")
    debut = decoder_texte(contenu[:4096]).lstrip()
    ext = _extension(nom_fichier)

    if debut.startswith("<"):
        return "xml"
    if debut.startswith(("{", "[")):
        try:
            json.loads(decoder_texte(contenu))
            return "json"
        except ValueError:
            pass
    if ext in ("csv", "tsv"):
        return "csv"
    if ext == "txt":
        return "txt"
    if ext in ("json", "xml"):
        raise DocumentIllisible(f"Le fichier .{ext} est invalide ou corrompu.")
    if debut and (debut.count(";") + debut.count("\t") + debut.count(",")) >= 3 and "\n" in debut:
        return "csv"
    return "txt"


# --------------------------------------------------------------------------
# OCR
# --------------------------------------------------------------------------

def _ocr_image(image) -> str:
    import pytesseract
    try:
        return pytesseract.image_to_string(image, lang=LANGUES_OCR)
    except pytesseract.TesseractError:
        # Une langue (ex. arabe) n'est pas installée : on retente avec l'anglais.
        return pytesseract.image_to_string(image, lang="eng")


def _ocr_disponible() -> bool:
    try:
        import pytesseract
        pytesseract.get_tesseract_version()
        return True
    except Exception:
        return False


# --------------------------------------------------------------------------
# Lecteurs par format
# --------------------------------------------------------------------------

def _lire_pdf(contenu: bytes) -> DocumentLu:
    from pypdf import PdfReader

    try:
        lecteur = PdfReader(io.BytesIO(contenu))
        if lecteur.is_encrypted:
            try:
                lecteur.decrypt("")
            except Exception as exc:
                raise DocumentIllisible("PDF protégé par mot de passe : impossible de le lire.") from exc
        textes = [(page.extract_text() or "").strip() for page in lecteur.pages]
    except DocumentIllisible:
        raise
    except Exception as exc:
        raise DocumentIllisible("PDF corrompu ou illisible.") from exc

    nb_pages = len(textes)
    doc = DocumentLu(format="pdf", texte="\n\n".join(t for t in textes if t), pages=nb_pages)

    moyenne = (sum(len(t) for t in textes) / nb_pages) if nb_pages else 0
    if nb_pages and moyenne < SEUIL_TEXTE_PAR_PAGE:
        doc.avertissements.append("PDF numérisé (peu ou pas de texte) : lecture par OCR.")
        if not _ocr_disponible():
            doc.avertissements.append("OCR indisponible sur ce serveur : le contenu n'a pas pu être lu.")
        else:
            try:
                import pypdfium2 as pdfium

                pages_ocr = []
                pdf = pdfium.PdfDocument(contenu)
                try:
                    for i in range(len(pdf)):
                        if i >= MAX_PAGES_OCR:
                            doc.avertissements.append(f"OCR limité aux {MAX_PAGES_OCR} premières pages.")
                            break
                        image = pdf[i].render(scale=200 / 72).to_pil()
                        pages_ocr.append(_ocr_image(image).strip())
                finally:
                    pdf.close()
                texte_ocr = "\n\n".join(t for t in pages_ocr if t)
                if texte_ocr:
                    doc.texte = texte_ocr
                    doc.ocr_utilise = True
            except Exception:
                doc.avertissements.append("Échec de l'OCR sur ce PDF.")
    return doc


def _lire_image(contenu: bytes) -> DocumentLu:
    from PIL import Image, ImageSequence

    try:
        image = Image.open(io.BytesIO(contenu))
        image.load()
    except Exception as exc:
        raise DocumentIllisible("Image corrompue ou illisible.") from exc
    if image.width * image.height > MAX_PIXELS_IMAGE:
        raise FichierTropVolumineux("Image trop grande (résolution excessive).")

    doc = DocumentLu(format="image", pages=getattr(image, "n_frames", 1))
    if not _ocr_disponible():
        doc.avertissements.append("OCR indisponible sur ce serveur : l'image n'a pas pu être lue.")
        return doc

    textes = []
    for i, cadre in enumerate(ImageSequence.Iterator(image)):
        if i >= MAX_PAGES_OCR:
            doc.avertissements.append(f"OCR limité aux {MAX_PAGES_OCR} premières pages.")
            break
        textes.append(_ocr_image(cadre.convert("L")).strip())
    doc.texte = "\n\n".join(t for t in textes if t)
    doc.ocr_utilise = True
    return doc


def _lire_docx(contenu: bytes) -> DocumentLu:
    from docx import Document

    try:
        document = Document(io.BytesIO(contenu))
    except Exception as exc:
        raise DocumentIllisible("Document Word corrompu ou illisible.") from exc

    lignes = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    tableaux = []
    for tableau in document.tables:
        rangees = [[cellule.text.strip() for cellule in rangee.cells] for rangee in tableau.rows]
        tableaux.append(rangees)
        lignes.extend(" | ".join(r) for r in rangees if any(r))
    return DocumentLu(format="docx", texte="\n".join(lignes), tableaux=tableaux)


def _texte_cellule(valeur) -> str:
    if valeur is None:
        return ""
    if hasattr(valeur, "isoformat"):
        return valeur.isoformat()
    return str(valeur).strip()


def _lire_xlsx(contenu: bytes) -> DocumentLu:
    from openpyxl import load_workbook

    try:
        classeur = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    except Exception as exc:
        raise DocumentIllisible("Classeur Excel corrompu ou illisible.") from exc

    doc = DocumentLu(format="xlsx")
    lignes_texte = []
    for feuille in classeur.worksheets:
        rangees = []
        for i, rangee in enumerate(feuille.iter_rows(values_only=True)):
            if i >= MAX_LIGNES_TABLEUR:
                doc.avertissements.append(f"Feuille « {feuille.title} » tronquée à {MAX_LIGNES_TABLEUR} lignes.")
                break
            cellules = [_texte_cellule(c) for c in rangee]
            if any(cellules):
                rangees.append(cellules)
        if rangees:
            doc.tableaux.append(rangees)
            lignes_texte.append(f"## Feuille : {feuille.title}")
            lignes_texte.extend(" | ".join(r) for r in rangees)
    classeur.close()
    doc.texte = "\n".join(lignes_texte)
    return doc


def _lire_csv(contenu: bytes) -> DocumentLu:
    texte = decoder_texte(contenu)
    echantillon = texte[:4096]
    try:
        dialecte = csv.Sniffer().sniff(echantillon, delimiters=",;\t|")
    except csv.Error:
        premiere = echantillon.splitlines()[0] if echantillon else ""
        dialecte = csv.excel_tab if "\t" in premiere else csv.excel
        if premiere.count(";") > premiere.count(","):
            dialecte = type("PointVirgule", (csv.excel,), {"delimiter": ";"})

    doc = DocumentLu(format="csv")
    rangees = []
    for i, rangee in enumerate(csv.reader(io.StringIO(texte), dialecte)):
        if i >= MAX_LIGNES_TABLEUR:
            doc.avertissements.append(f"CSV tronqué à {MAX_LIGNES_TABLEUR} lignes.")
            break
        cellules = [c.strip() for c in rangee]
        if any(cellules):
            rangees.append(cellules)
    doc.tableaux = [rangees] if rangees else []
    doc.texte = "\n".join(" | ".join(r) for r in rangees)
    return doc


def _parcourir_formes(formes, lignes, tableaux):
    for forme in formes:
        if getattr(forme, "has_text_frame", False) and forme.text_frame.text.strip():
            lignes.append(forme.text_frame.text.strip())
        if getattr(forme, "has_table", False) and forme.has_table:
            rangees = [[c.text.strip() for c in r.cells] for r in forme.table.rows]
            tableaux.append(rangees)
            lignes.extend(" | ".join(r) for r in rangees if any(r))
        if hasattr(forme, "shapes"):          # groupe de formes
            _parcourir_formes(forme.shapes, lignes, tableaux)


def _lire_pptx(contenu: bytes) -> DocumentLu:
    from pptx import Presentation

    try:
        presentation = Presentation(io.BytesIO(contenu))
    except Exception as exc:
        raise DocumentIllisible("Présentation PowerPoint corrompue ou illisible.") from exc

    lignes, tableaux = [], []
    for numero, diapo in enumerate(presentation.slides, start=1):
        lignes.append(f"## Diapositive {numero}")
        _parcourir_formes(diapo.shapes, lignes, tableaux)
    return DocumentLu(format="pptx", texte="\n".join(lignes), tableaux=tableaux, pages=len(presentation.slides))


def _nom_sans_espace_de_noms(balise: str) -> str:
    return balise.rsplit("}", 1)[-1]


def _lire_xml(contenu: bytes) -> DocumentLu:
    from defusedxml import ElementTree
    from defusedxml.common import DefusedXmlException

    try:
        racine = ElementTree.fromstring(contenu)
    except DefusedXmlException as exc:
        raise DocumentIllisible("XML refusé : entités ou DTD externes interdites (sécurité).") from exc
    except ElementTree.ParseError as exc:
        raise DocumentIllisible("Fichier XML invalide.") from exc

    lignes = []

    def descendre(element, chemin):
        nom = f"{chemin}/{_nom_sans_espace_de_noms(element.tag)}"
        for attribut, valeur in element.attrib.items():
            lignes.append(f"{nom}@{_nom_sans_espace_de_noms(attribut)}: {valeur}")
        enfants = list(element)
        valeur = (element.text or "").strip()
        if valeur and not enfants:
            lignes.append(f"{nom}: {valeur}")
        for enfant in enfants:
            if len(lignes) < MAX_CARACTERES // 20:
                descendre(enfant, nom)

    descendre(racine, "")
    return DocumentLu(format="xml", texte="\n".join(lignes))


def _lire_json(contenu: bytes) -> DocumentLu:
    try:
        donnees = json.loads(decoder_texte(contenu))
    except ValueError as exc:
        raise DocumentIllisible("Fichier JSON invalide.") from exc
    return DocumentLu(format="json", texte=json.dumps(donnees, indent=2, ensure_ascii=False))


def _lire_txt(contenu: bytes) -> DocumentLu:
    return DocumentLu(format="txt", texte=decoder_texte(contenu).strip())


_LECTEURS = {
    "pdf": _lire_pdf, "image": _lire_image, "docx": _lire_docx, "xlsx": _lire_xlsx,
    "csv": _lire_csv, "pptx": _lire_pptx, "xml": _lire_xml, "json": _lire_json, "txt": _lire_txt,
}


def lire_document(contenu: bytes, nom_fichier: str = "") -> DocumentLu:
    """Point d'entrée : détecte le format par le contenu puis extrait texte et tableaux."""
    if len(contenu) > TAILLE_MAX_OCTETS:
        raise FichierTropVolumineux(
            f"Fichier trop volumineux (maximum {TAILLE_MAX_OCTETS // (1024 * 1024)} Mo)."
        )
    format_detecte = detecter_format(contenu, nom_fichier)
    doc = _LECTEURS[format_detecte](contenu)

    if len(doc.texte) > MAX_CARACTERES:
        doc.texte = doc.texte[:MAX_CARACTERES]
        doc.avertissements.append("Texte tronqué : document très long.")
    if not doc.texte.strip() and not doc.tableaux:
        doc.avertissements.append("Aucun texte détecté dans le document.")
    return doc
