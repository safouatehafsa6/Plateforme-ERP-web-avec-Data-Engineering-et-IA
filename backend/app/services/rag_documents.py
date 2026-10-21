"""Import de documents pour le RAG : validation, extraction du texte, découpe
en chunks et stockage de l'original dans MinIO.

Limite connue : pas d'OCR. Un PDF scanné (images sans couche texte) est refusé
avec un message explicite.
"""
import io
import logging
import os
import re

from minio import Minio
from minio.error import S3Error

logger = logging.getLogger("erp.rag")

BUCKET_RAG = os.getenv("RAG_BUCKET", "rag-documents")
TAILLE_MAX_OCTETS = 10 * 1024 * 1024  # 10 Mo
TYPES_AUTORISES = {"pdf", "docx"}
CHUNK_TAILLE = 900
CHUNK_CHEVAUCHEMENT = 150
MAX_CHUNKS = 2000  # garde-fou contre les documents démesurés


class DocumentInvalide(Exception):
    """Erreur imputable au fichier : message affichable à l'utilisateur."""


# --------------------------------------------------------------------- MinIO
_client_minio = None


def _client() -> Minio:
    global _client_minio
    if _client_minio is None:
        _client_minio = Minio(
            os.environ["MINIO_ENDPOINT"],
            access_key=os.environ["MINIO_ROOT_USER"],
            secret_key=os.environ["MINIO_ROOT_PASSWORD"],
            secure=os.getenv("MINIO_SECURE", "false").lower() == "true",
        )
    return _client_minio


def assurer_bucket() -> None:
    client = _client()
    if not client.bucket_exists(BUCKET_RAG):
        try:
            client.make_bucket(BUCKET_RAG)
        except S3Error as e:  # course possible entre deux requêtes
            if e.code not in ("BucketAlreadyOwnedByYou", "BucketAlreadyExists"):
                raise


def stocker_original(cle: str, contenu: bytes, type_mime: str) -> None:
    assurer_bucket()
    _client().put_object(
        BUCKET_RAG, cle, io.BytesIO(contenu), length=len(contenu), content_type=type_mime
    )


def supprimer_original(cle: str) -> None:
    try:
        _client().remove_object(BUCKET_RAG, cle)
    except Exception:  # ne jamais masquer l'erreur principale
        logger.exception("Suppression MinIO impossible pour %s", cle)


# ---------------------------------------------------------------- Validation
def nettoyer_nom_fichier(nom: str | None) -> str:
    """Nom d'affichage seulement : il n'est JAMAIS utilisé comme chemin MinIO."""
    base = os.path.basename((nom or "").replace("\\", "/"))
    base = re.sub(r"[\x00-\x1f]", "", base).strip()
    return (base or "document")[:200]


def valider_fichier(nom: str, contenu: bytes) -> str:
    """Retourne le type ('pdf'|'docx') ou lève DocumentInvalide."""
    ext = nom.rsplit(".", 1)[-1].lower() if "." in nom else ""
    if ext not in TYPES_AUTORISES:
        raise DocumentInvalide("Format non autorisé : seuls les fichiers PDF et DOCX sont acceptés.")
    if not contenu:
        raise DocumentInvalide("Le fichier est vide.")
    # On vérifie la signature réelle, pas seulement l'extension.
    if ext == "pdf" and not contenu.lstrip()[:5] == b"%PDF-":
        raise DocumentInvalide("Le fichier n'est pas un PDF valide.")
    if ext == "docx" and not contenu.startswith(b"PK"):
        raise DocumentInvalide("Le fichier n'est pas un DOCX valide.")
    return ext


# ---------------------------------------------------------------- Extraction
def extraire_pages(contenu: bytes, type_fichier: str) -> list[tuple[int | None, str]]:
    """Retourne [(numéro_de_page|None, texte)]."""
    try:
        if type_fichier == "pdf":
            from pypdf import PdfReader

            lecteur = PdfReader(io.BytesIO(contenu))
            if lecteur.is_encrypted:
                raise DocumentInvalide("Le PDF est protégé par mot de passe.")
            pages = []
            for i, page in enumerate(lecteur.pages, 1):
                texte = (page.extract_text() or "").strip()
                if texte:
                    pages.append((i, texte))
            return pages

        from docx import Document

        doc = Document(io.BytesIO(contenu))
        morceaux = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for ligne in table.rows:
                cellules = [c.text.strip() for c in ligne.cells if c.text.strip()]
                if cellules:
                    morceaux.append(" | ".join(cellules))
        texte = "\n".join(morceaux).strip()
        return [(None, texte)] if texte else []
    except DocumentInvalide:
        raise
    except Exception:
        logger.exception("Échec d'extraction du texte (%s)", type_fichier)
        raise DocumentInvalide("Impossible de lire ce fichier : il est peut-être corrompu.")


# ------------------------------------------------------------------ Chunking
def decouper_texte(texte: str, taille: int = CHUNK_TAILLE, chevauchement: int = CHUNK_CHEVAUCHEMENT) -> list[str]:
    texte = re.sub(r"[ \t]+", " ", texte.replace("\x00", ""))
    texte = re.sub(r"\n{3,}", "\n\n", texte).strip()
    if not texte:
        return []
    chunks, debut = [], 0
    while debut < len(texte):
        fin = min(debut + taille, len(texte))
        if fin < len(texte):
            # coupe de préférence sur un saut de ligne, une fin de phrase ou un espace
            fenetre = texte[debut:fin]
            coupe = max(fenetre.rfind("\n"), fenetre.rfind(". "), fenetre.rfind(" "))
            if coupe > taille * 0.5:
                fin = debut + coupe + 1
        morceau = texte[debut:fin].strip()
        if morceau:
            chunks.append(morceau)
        if fin >= len(texte):
            break
        debut = max(fin - chevauchement, debut + 1)
    return chunks


def construire_chunks(pages, nom_fichier: str) -> list[tuple[str, str]]:
    """[(source, contenu)] — la source indique le fichier et la page."""
    resultat = []
    for numero, texte in pages:
        for morceau in decouper_texte(texte):
            source = f"{nom_fichier} (p.{numero})" if numero else nom_fichier
            resultat.append((source, morceau))
    if len(resultat) > MAX_CHUNKS:
        raise DocumentInvalide("Document trop volumineux à indexer (trop de contenu textuel).")
    return resultat
