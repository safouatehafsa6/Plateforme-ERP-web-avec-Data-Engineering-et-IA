"""Rapports PDF et Word des calculs (taux unique, taux variables détaillé / simplifié)."""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from app.services.moteur_calcul import MODULES_SOURCE

MODES = {"taux_unique": "Calcul à taux unique", "taux_variable": "Calcul à taux variables"}
TYPES_RAPPORT = ("detaille", "simplifie")


def nombre_fr(valeur, decimales: int = 2) -> str:
    """12345.6 -> « 12 345,60 » (espace insécable comme séparateur de milliers)."""
    texte = f"{Decimal(str(valeur)):,.{decimales}f}"
    return texte.replace(",", "\u00a0").replace(".", ",")


def date_fr(valeur) -> str:
    if isinstance(valeur, str):
        valeur = date.fromisoformat(valeur[:10])
    return valeur.strftime("%d/%m/%Y")


def _taux_fr(valeur) -> str:
    return nombre_fr(valeur, 2) + " %"


def _lignes_infos(c: dict) -> list[tuple[str, str]]:
    devise = c["devise"]
    infos = [
        ("Référence", c["reference"]),
        ("Type de calcul", MODES.get(c["mode"], c["mode"])),
        ("Module concerné", MODULES_SOURCE.get(c["module_source"], c["module_source"])),
    ]
    if c.get("libelle"):
        infos.append(("Objet", c["libelle"]))
    infos += [
        ("Montant", f"{nombre_fr(c['montant'])} {devise}"),
        ("Période", f"du {date_fr(c['date_debut'])} au {date_fr(c['date_fin'])}"),
        ("Nombre de jours", str(c["nombre_jours"])),
    ]
    if c["mode"] == "taux_unique":
        infos.append(("Taux", _taux_fr(c["taux"])))
    elif c.get("bareme_libelle"):
        infos.append(("Barème de taux", c["bareme_libelle"]))
    infos.append(("Établi par", f"{c.get('cree_par_nom') or c['cree_par']} le {date_fr(c['cree_le'])}"))
    return infos


def _formule(c: dict) -> str:
    return (f"Résultat = Montant × (Taux / 100) × (Nombre de jours / {c['base_jours']})  =  "
            f"{nombre_fr(c['montant'])} × ({nombre_fr(c['taux'])} / 100) × ({c['nombre_jours']} / {c['base_jours']})")


def _tableau_periodes(c: dict) -> list[list[str]]:
    devise = c["devise"]
    lignes = [["Période", "Jours", "Taux", f"Résultat ({devise})"]]
    for d in c["details"]:
        lignes.append([f"{date_fr(d['debut'])} → {date_fr(d['fin'])}", str(d["jours"]), _taux_fr(d["taux"]), nombre_fr(d["resultat"])])
    lignes.append(["Total", str(c["nombre_jours"]), "", nombre_fr(c["resultat"])])
    return lignes


def _titre(c: dict, type_rapport: str) -> str:
    if c["mode"] == "taux_unique":
        return "Rapport de calcul — taux unique"
    return "Rapport de calcul — taux variables (" + ("détaillé" if type_rapport == "detaille" else "simplifié") + ")"


def generer_pdf(c: dict, entreprise: str, type_rapport: str = "detaille") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    titre = ParagraphStyle("t", parent=styles["Title"], fontSize=17, spaceAfter=4)
    sous = ParagraphStyle("s", parent=styles["Normal"], textColor=colors.HexColor("#555555"), alignment=1, spaceAfter=14)
    petit = ParagraphStyle("p", parent=styles["Normal"], fontSize=8.5, textColor=colors.HexColor("#555555"))
    corps = ParagraphStyle("c", parent=styles["Normal"], fontSize=10)

    tampon = BytesIO()
    doc = SimpleDocTemplate(tampon, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                            title=f"{c['reference']} - {_titre(c, type_rapport)}")
    histoire = [Paragraph(_titre(c, type_rapport), titre), Paragraph(entreprise, sous)]

    infos = Table([[Paragraph(f"<b>{k}</b>", corps), Paragraph(v, corps)] for k, v in _lignes_infos(c)], colWidths=[5 * cm, 12 * cm])
    infos.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("LINEBELOW", (0, 0), (-1, -1), 0.4, colors.HexColor("#DDDDDD")),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    histoire += [infos, Spacer(1, 14)]

    if c["mode"] == "taux_variable" and type_rapport == "detaille":
        t = Table(_tableau_periodes(c), colWidths=[7.2 * cm, 2.2 * cm, 3 * cm, 4.6 * cm], repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#14283C")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"),
            ("ALIGN", (1, 0), (-1, -1), "RIGHT"), ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -2), [colors.white, colors.HexColor("#F5F3EF")]),
            ("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#EBDDBF")), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
        ]))
        histoire.append(t)
    else:
        if c["mode"] == "taux_unique":
            histoire += [Paragraph(_formule(c), petit), Spacer(1, 10)]
        resultat = Table([[Paragraph("<b>Résultat</b>", corps), Paragraph(f"<b>{nombre_fr(c['resultat'])} {c['devise']}</b>", ParagraphStyle("r", parent=corps, fontSize=14, alignment=2))]], colWidths=[5 * cm, 12 * cm])
        resultat.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EBDDBF")), ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 10)]))
        histoire.append(resultat)

    histoire += [Spacer(1, 18), Paragraph(
        f"Base de calcul : {c['base_jours']} jours par an. Les montants sont arrondis au centime. "
        f"Document généré le {datetime.now():%d/%m/%Y à %H:%M}.", petit)]
    doc.build(histoire)
    return tampon.getvalue()


def generer_docx(c: dict, entreprise: str, type_rapport: str = "detaille") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    document = Document()
    titre = document.add_heading(_titre(c, type_rapport), level=1)
    titre.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sous = document.add_paragraph(entreprise)
    sous.alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = document.add_table(rows=0, cols=2)
    table.style = "Light List Accent 1"
    for cle, valeur in _lignes_infos(c):
        cellules = table.add_row().cells
        cellules[0].text = cle
        cellules[1].text = valeur
        cellules[0].paragraphs[0].runs[0].bold = True
    document.add_paragraph()

    if c["mode"] == "taux_variable" and type_rapport == "detaille":
        lignes = _tableau_periodes(c)
        t = document.add_table(rows=len(lignes), cols=4)
        t.style = "Light Grid Accent 1"
        for i, ligne in enumerate(lignes):
            for j, texte in enumerate(ligne):
                cellule = t.cell(i, j)
                cellule.text = texte
                if j > 0:
                    cellule.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
                if i == 0 or i == len(lignes) - 1:
                    for run in cellule.paragraphs[0].runs:
                        run.bold = True
    else:
        if c["mode"] == "taux_unique":
            p = document.add_paragraph(_formule(c))
            p.runs[0].font.size = Pt(9)
        p = document.add_paragraph()
        run = p.add_run(f"Résultat : {nombre_fr(c['resultat'])} {c['devise']}")
        run.bold = True
        run.font.size = Pt(15)

    note = document.add_paragraph(
        f"Base de calcul : {c['base_jours']} jours par an. Les montants sont arrondis au centime. "
        f"Document généré le {datetime.now():%d/%m/%Y à %H:%M}.")
    note.runs[0].font.size = Pt(8)
    tampon = BytesIO()
    document.save(tampon)
    return tampon.getvalue()
