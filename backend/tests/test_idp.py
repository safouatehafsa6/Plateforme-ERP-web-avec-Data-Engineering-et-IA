"""Tests de l'IDP (étape 1) : lecture des formats et contrôles de cohérence.

Lancer depuis le dossier backend :
    python -m unittest tests.test_idp -v
(dans Docker : docker compose exec backend python -m unittest tests.test_idp -v)
"""
import io
import json
import shutil
import unittest
from datetime import date
from decimal import Decimal
from unittest import mock

from app.services import idp_extraction as ext
from app.services.idp_controles import (
    a_des_erreurs, convertir_date, convertir_decimal, empreinte_document,
    resume_anomalies, verifier_document,
)

AUJOURDHUI = date(2026, 10, 26)


def facture_valide():
    return {
        "type_document": "facture", "numero": "F-2026-0042", "date": "2026-10-12",
        "date_echeance": "2026-11-12", "devise": "MAD",
        "fournisseur": {"nom": "Atlas Fournitures SARL", "ice": "001234567000012"},
        "lignes": [
            {"designation": "Papier A4", "quantite": 10, "prix_unitaire": "25,00", "taux_tva": 20, "total_ht": 250},
            {"designation": "Stylos", "quantite": 5, "prix_unitaire": 10, "remise_pct": 10, "taux_tva": 20},
        ],
        "total_ht": "295,00", "total_tva": "59,00", "total_ttc": "354,00",
    }


def codes(anomalies):
    return {a.code for a in anomalies}


class TestConversions(unittest.TestCase):
    def test_decimaux(self):
        cas = {
            "1 234,56": "1234.56", "1.234,56": "1234.56", "1,234.56": "1234.56",
            "1 234,56 MAD": "1234.56", "(50,00)": "-50.00", "-12,5": "-12.5",
            "1.234.567": "1234567", "12": "12", 12.5: "12.5",
        }
        for entree, attendu in cas.items():
            self.assertEqual(convertir_decimal(entree), Decimal(attendu), entree)
        for invalide in (None, "", "abc", True):
            self.assertIsNone(convertir_decimal(invalide), invalide)

    def test_dates(self):
        for texte in ("2026-10-12", "12/10/2026", "12-10-2026", "12.10.2026"):
            self.assertEqual(convertir_date(texte), date(2026, 10, 12), texte)
        self.assertIsNone(convertir_date("pas une date"))
        self.assertIsNone(convertir_date(None))


class TestControles(unittest.TestCase):
    def verifier(self, donnees):
        return verifier_document(donnees, aujourdhui=AUJOURDHUI)

    def test_facture_coherente(self):
        anomalies = self.verifier(facture_valide())
        self.assertEqual(anomalies, [], [a.vers_dict() for a in anomalies])
        self.assertFalse(a_des_erreurs(anomalies))

    def test_total_ligne_faux(self):
        d = facture_valide()
        d["lignes"][0]["total_ht"] = 300
        self.assertIn("total_ligne_incoherent", codes(self.verifier(d)))

    def test_somme_ht_fausse(self):
        d = facture_valide()
        d["total_ht"] = "400,00"
        self.assertIn("total_ht_incoherent", codes(self.verifier(d)))

    def test_tva_fausse(self):
        d = facture_valide()
        d["total_tva"] = "30,00"
        d["total_ttc"] = "325,00"
        self.assertIn("total_tva_incoherent", codes(self.verifier(d)))

    def test_ttc_different_de_ht_plus_tva(self):
        d = facture_valide()
        d["total_ttc"] = "400,00"
        self.assertIn("total_ttc_incoherent", codes(self.verifier(d)))

    def test_champs_manquants(self):
        trouves = codes(self.verifier({}))
        for attendu in ("numero_manquant", "fournisseur_manquant", "date_manquante",
                        "lignes_manquantes", "total_ttc_manquant"):
            self.assertIn(attendu, trouves)

    def test_dates_incoherentes(self):
        d = facture_valide()
        d["date_echeance"] = "2026-09-01"
        self.assertIn("echeance_avant_date", codes(self.verifier(d)))
        d = facture_valide()
        d["date"] = "2027-05-01"
        self.assertIn("date_future", codes(self.verifier(d)))
        d["date"] = "31/02/2026"
        self.assertIn("date_invalide", codes(self.verifier(d)))

    def test_quantite_et_remise_invalides(self):
        d = facture_valide()
        d["lignes"][0]["quantite"] = 0
        d["lignes"][1]["remise_pct"] = 150
        trouves = codes(self.verifier(d))
        self.assertIn("quantite_invalide", trouves)
        self.assertIn("remise_invalide", trouves)

    def test_avertissements_non_bloquants(self):
        d = facture_valide()
        d["devise"] = "XYZ"
        d["fournisseur"]["ice"] = "123"
        d["lignes"][0]["taux_tva"] = 13
        d["total_tva"] = None
        anomalies = self.verifier(d)
        self.assertTrue({"devise_inconnue", "ice_invalide", "taux_tva_inhabituel"} <= codes(anomalies))
        self.assertFalse(a_des_erreurs([a for a in anomalies if a.code in
                                        {"devise_inconnue", "ice_invalide", "taux_tva_inhabituel"}]))

    def test_avoir_accepte_les_montants_negatifs(self):
        d = facture_valide()
        d["type_document"] = "avoir"
        d["lignes"] = [{"designation": "Retour", "quantite": -2, "prix_unitaire": 10, "taux_tva": 20}]
        d["total_ht"], d["total_tva"], d["total_ttc"] = "-20", "-4", "-24"
        self.assertEqual(self.verifier(d), [])

    def test_resume(self):
        d = facture_valide()
        d["total_ht"] = "999"
        resume = resume_anomalies(self.verifier(d))
        self.assertGreaterEqual(resume["erreurs"], 1)

    def test_empreinte_doublon(self):
        a, b = facture_valide(), facture_valide()
        b["fournisseur"]["nom"] = "  ATLAS   fournitures sarl "
        self.assertEqual(empreinte_document(a), empreinte_document(b))
        b["total_ttc"] = "355,00"
        self.assertNotEqual(empreinte_document(a), empreinte_document(b))
        self.assertIsNone(empreinte_document({"lignes": []}))


class TestExtraction(unittest.TestCase):
    def test_csv_point_virgule(self):
        doc = ext.lire_document("ref;désignation;prix\nA1;Papier;25,00\nA2;Stylo;10,00\n".encode("utf-8"), "x.csv")
        self.assertEqual(doc.format, "csv")
        self.assertEqual(doc.tableaux[0][1], ["A1", "Papier", "25,00"])

    def test_csv_encodage_windows(self):
        doc = ext.lire_document("nom;ville\nÉlodie;Fès\n".encode("cp1252"), "x.csv")
        self.assertIn("Élodie", doc.texte)

    def test_json_et_txt(self):
        doc = ext.lire_document(json.dumps({"numero": "F1", "total": 10}).encode(), "a.json")
        self.assertEqual(doc.format, "json")
        self.assertIn('"numero": "F1"', doc.texte)
        doc = ext.lire_document("Bonjour, ceci est une note.".encode(), "n.txt")
        self.assertEqual(doc.format, "txt")

    def test_xml(self):
        xml = b'<?xml version="1.0"?><facture numero="F1"><fournisseur>Atlas</fournisseur><total>10</total></facture>'
        doc = ext.lire_document(xml, "f.xml")
        self.assertEqual(doc.format, "xml")
        self.assertIn("/facture/fournisseur: Atlas", doc.texte)
        self.assertIn("/facture@numero: F1", doc.texte)

    def test_xml_xxe_refuse(self):
        piege = (b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY x SYSTEM "file:///etc/passwd">]>'
                 b'<d>&x;</d>')
        with self.assertRaises(ext.DocumentIllisible):
            ext.lire_document(piege, "f.xml")

    def test_docx(self):
        from docx import Document
        d = Document()
        d.add_paragraph("Facture F-2026-0042")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Article", "Prix"
        t.cell(1, 0).text, t.cell(1, 1).text = "Papier", "25,00"
        tampon = io.BytesIO()
        d.save(tampon)
        doc = ext.lire_document(tampon.getvalue(), "f.docx")
        self.assertEqual(doc.format, "docx")
        self.assertIn("F-2026-0042", doc.texte)
        self.assertEqual(doc.tableaux[0][1], ["Papier", "25,00"])

    def test_xlsx(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "Lignes"
        ws.append(["Article", "Quantité", "Prix"])
        ws.append(["Papier", 10, 25.5])
        tampon = io.BytesIO()
        wb.save(tampon)
        doc = ext.lire_document(tampon.getvalue(), "f.xlsx")
        self.assertEqual(doc.format, "xlsx")
        self.assertEqual(doc.tableaux[0][1], ["Papier", "10", "25.5"])

    def test_pptx(self):
        from pptx import Presentation
        p = Presentation()
        diapo = p.slides.add_slide(p.slide_layouts[5])
        diapo.shapes.title.text = "Bon de commande BC-77"
        tampon = io.BytesIO()
        p.save(tampon)
        doc = ext.lire_document(tampon.getvalue(), "f.pptx")
        self.assertEqual(doc.format, "pptx")
        self.assertIn("BC-77", doc.texte)

    def _pdf_texte(self):
        from reportlab.pdfgen import canvas
        tampon = io.BytesIO()
        c = canvas.Canvas(tampon)
        c.drawString(72, 750, "Facture F-2026-0042 - Atlas Fournitures - Total TTC 354,00 MAD")
        c.save()
        return tampon.getvalue()

    def test_pdf_texte(self):
        doc = ext.lire_document(self._pdf_texte(), "f.pdf")
        self.assertEqual(doc.format, "pdf")
        self.assertFalse(doc.ocr_utilise)
        self.assertIn("F-2026-0042", doc.texte)

    def test_pdf_vide_est_signale(self):
        from reportlab.pdfgen import canvas
        tampon = io.BytesIO()
        c = canvas.Canvas(tampon)
        c.showPage()
        c.save()
        doc = ext.lire_document(tampon.getvalue(), "scan.pdf")
        self.assertTrue(doc.avertissements)

    @unittest.skipUnless(shutil.which("tesseract"), "tesseract non installé")
    def test_ocr_image(self):
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new("RGB", (900, 200), "white")
        ImageDraw.Draw(image).text((30, 60), "FACTURE 12345", fill="black", font=ImageFont.load_default(size=60))
        tampon = io.BytesIO()
        image.save(tampon, format="PNG")
        doc = ext.lire_document(tampon.getvalue(), "scan.png")
        self.assertEqual(doc.format, "image")
        self.assertTrue(doc.ocr_utilise)
        self.assertIn("12345", doc.texte)

    @unittest.skipUnless(shutil.which("tesseract"), "tesseract non installé")
    def test_ocr_pdf_numerise(self):
        from PIL import Image, ImageDraw, ImageFont
        image = Image.new("RGB", (1000, 300), "white")
        ImageDraw.Draw(image).text((30, 100), "BON DE COMMANDE 98765", fill="black", font=ImageFont.load_default(size=60))
        tampon = io.BytesIO()
        image.save(tampon, format="PDF")          # PDF contenant uniquement une image
        doc = ext.lire_document(tampon.getvalue(), "scan.pdf")
        self.assertEqual(doc.format, "pdf")
        self.assertTrue(doc.ocr_utilise)
        self.assertIn("98765", doc.texte)

    def test_extension_mensongere(self):
        # Un PDF déguisé en .docx est lu pour ce qu'il est.
        self.assertEqual(ext.lire_document(self._pdf_texte(), "pas_un_word.docx").format, "pdf")

    def test_ancien_format_office_refuse(self):
        faux_doc = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 100
        with self.assertRaises(ext.FormatNonSupporte):
            ext.lire_document(faux_doc, "vieux.doc")

    def test_fichier_vide_ou_inconnu(self):
        with self.assertRaises(ext.DocumentIllisible):
            ext.lire_document(b"", "x.pdf")
        with self.assertRaises(ext.DocumentIllisible):
            ext.lire_document(b"\x00\x01\x02\x03binaire", "x.bin")

    def test_fichier_trop_gros(self):
        with mock.patch.object(ext, "TAILLE_MAX_OCTETS", 10):
            with self.assertRaises(ext.FichierTropVolumineux):
                ext.lire_document(b"x" * 11, "x.txt")

    def test_bombe_de_decompression_refusee(self):
        from docx import Document
        tampon = io.BytesIO()
        Document().save(tampon)
        with mock.patch.object(ext, "TAILLE_MAX_DECOMPRESSEE", 1000):
            with self.assertRaises(ext.FichierTropVolumineux):
                ext.lire_document(tampon.getvalue(), "f.docx")

    def test_empreinte_fichier(self):
        self.assertEqual(ext.empreinte_fichier(b"abc"), ext.empreinte_fichier(b"abc"))
        self.assertNotEqual(ext.empreinte_fichier(b"abc"), ext.empreinte_fichier(b"abd"))


if __name__ == "__main__":
    unittest.main()
